import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference

# ------------------------------
# Step 1: Load the given Excel file
# ------------------------------
# Read directly from the "Student Marks" sheet inside student.xlsx
input_file = "student.xlsx"
df = pd.read_excel(input_file, sheet_name="Student Marks")

# ------------------------------
# Step 2: Vectorized Computations (no loops)
# ------------------------------
# Calculate total marks across subjects for each student
df["Total"] = df[["Math", "Physics", "Chemistry", "Biology"]].sum(axis=1) 

# Calculate average marks for each student
df["Average"] = df[["Math", "Physics", "Chemistry", "Biology"]].mean(axis=1)

# Assign grades using conditions
# We use np.select to avoid writing loops or multiple if-else statements
conditions = [
    df["Average"] >= 90,
    (df["Average"] >= 75) & (df["Average"] < 90),
    (df["Average"] >= 60) & (df["Average"] < 75),
    df["Average"] < 60
]
grades = ["A", "B", "C", "F"]
df["Grade"] = np.select(conditions, grades)


# Step 3: Find Top Performers per subject

subjects = ["Math", "Physics", "Chemistry", "Biology"]

# Create an empty DataFrame to hold results
top_df = pd.DataFrame()

for subject in subjects:
    # Take top 3 students directly
    top3 = df.sort_values(by=subject, ascending=False).head(3)
    
    # Add a new column to indicate subject
    top3 = top3.assign(Subject=subject, Marks=top3[subject])
    
    # Keep only the useful columns
    top3 = top3[["Subject", "StudentID", "Name", "Marks"]]
    
    # Append directly into the result DataFrame
    top_df = pd.concat([top_df, top3], ignore_index=True)


# ------------------------------
# Step 4: Subject Averages DataFrame
# ------------------------------
subject_avg_df = (
    df[["Math", "Physics", "Chemistry", "Biology"]]
    .mean()
    .reset_index() # convert the series into dataframe [index,0]
    .rename(columns={"index": "Subject", 0: "Average Marks"})
)

# ------------------------------
# Step 5: Save results back into the same Excel
# ------------------------------
with pd.ExcelWriter(input_file, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
    df[["StudentID", "Name", "Total", "Average", "Grade"]].to_excel(writer, sheet_name="Summary", index=False)
    top_df.to_excel(writer, sheet_name="Top Performers", index=False)
    subject_avg_df.to_excel(writer, sheet_name="Subject Averages", index=False)

# ------------------------------
# Step 6: Add bar chart to Excel
# ------------------------------
wb = load_workbook(input_file)
ws = wb["Subject Averages"]

chart = BarChart()
data = Reference(ws, min_col=2, min_row=1, max_row=5)  # Average Marks column including header
cats = Reference(ws, min_col=1, min_row=2, max_row=5)  # Subjects (no header)

chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.title = "Average Marks per Subject"
chart.y_axis.title = "Marks"
chart.x_axis.title = "Subjects"

ws.add_chart(chart, "D2")  # Place chart at cell D2

wb.save(input_file)